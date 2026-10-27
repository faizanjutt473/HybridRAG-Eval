import os
import time
import requests
from docx import Document

OUT_DIR = "data/raw_docs"
os.makedirs(OUT_DIR, exist_ok=True)

HEADERS = {"User-Agent": "HybridRAG-Eval/1.0 (student project)"}

# ---------- 1) 10 PDFs from arXiv ----------
PDFS = {
    "attention_is_all_you_need.pdf": "https://arxiv.org/pdf/1706.03762",
    "bert.pdf": "https://arxiv.org/pdf/1810.04805",
    "rag_lewis_et_al.pdf": "https://arxiv.org/pdf/2005.11401",
    "resnet.pdf": "https://arxiv.org/pdf/1512.03385",
    "lora.pdf": "https://arxiv.org/pdf/2106.09685",
    "adam_optimizer.pdf": "https://arxiv.org/pdf/1412.6980",
    "batch_normalization.pdf": "https://arxiv.org/pdf/1502.03167",
    "word2vec.pdf": "https://arxiv.org/pdf/1301.3781",
    "dropout.pdf": "https://arxiv.org/pdf/1207.0580",
    "dense_passage_retrieval.pdf": "https://arxiv.org/pdf/2004.04906",
}

# ---------- 2) Wikipedia articles ----------
TXT_ARTICLES = {
    "wiki_machine_learning.txt": "Machine learning",
    "wiki_overfitting.txt": "Overfitting",
    "wiki_gradient_descent.txt": "Gradient descent",
    "wiki_random_forest.txt": "Random forest",
    "wiki_convolutional_neural_network.txt": "Convolutional neural network",
    "wiki_reinforcement_learning.txt": "Reinforcement learning",
}

DOCX_ARTICLES = {
    "wiki_transformer_model.docx": "Transformer (deep learning architecture)",
    "wiki_support_vector_machine.docx": "Support vector machine",
    "wiki_k_means_clustering.docx": "K-means clustering",
    "wiki_backpropagation.docx": "Backpropagation",
    "wiki_word_embedding.docx": "Word embedding",
    "wiki_recurrent_neural_network.docx": "Recurrent neural network",
}


def get_wikipedia_text(title):
    r = requests.get(
        "https://en.wikipedia.org/w/api.php",
        params={
            "action": "query",
            "prop": "extracts",
            "explaintext": 1,
            "titles": title,
            "format": "json",
            "redirects": 1,
        },
        headers=HEADERS,
        timeout=60,
    )
    r.raise_for_status()
    pages = r.json()["query"]["pages"]
    return next(iter(pages.values())).get("extract", "")


# --- PDFs ---
print("Downloading PDFs...")
for name, url in PDFS.items():
    path = os.path.join(OUT_DIR, name)
    try:
        r = requests.get(url, headers=HEADERS, timeout=120)
        r.raise_for_status()
        with open(path, "wb") as f:
            f.write(r.content)
        print("  OK  ", name, f"({len(r.content)//1024} KB)")
    except Exception as e:
        print("  FAIL", name, e)
    time.sleep(3)  # arXiv ko overload na karne ke liye

# --- TXT ---
print("Creating TXT files...")
for name, title in TXT_ARTICLES.items():
    try:
        text = get_wikipedia_text(title)
        with open(os.path.join(OUT_DIR, name), "w", encoding="utf-8") as f:
            f.write(f"{title}\n\n{text}")
        print("  OK  ", name, f"({len(text)} chars)")
    except Exception as e:
        print("  FAIL", name, e)

# --- DOCX ---
print("Creating DOCX files...")
for name, title in DOCX_ARTICLES.items():
    try:
        text = get_wikipedia_text(title)
        doc = Document()
        doc.add_heading(title, level=0)
        for para in text.split("\n"):
            para = para.strip()
            if not para:
                continue
            if para.startswith("==") and para.endswith("=="):
                doc.add_heading(para.strip("= ").strip(), level=1)
            else:
                doc.add_paragraph(para)
        doc.save(os.path.join(OUT_DIR, name))
        print("  OK  ", name, f"({len(text)} chars)")
    except Exception as e:
        print("  FAIL", name, e)

print("\nDone. Files in", OUT_DIR, ":", len(os.listdir(OUT_DIR)))